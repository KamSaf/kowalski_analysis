import pandas as pd
from datetime import datetime
import os
from docx import Document
import tqdm as tqdm

INPUT_PATH = r"car_prices_cleansed.csv"
OUTPUT_PATH = r"car_prices_cleansed_temp.csv"


def generate_head_doc():
    # Przetwarzanie linia po linii
    with open(INPUT_PATH, "r") as infile, open(OUTPUT_PATH, "w") as outfile:
        for line in infile:
            # Zamiana ";" na ","
            outfile.write(line.replace(";", ","))

    # Zastąpienie oryginalnego pliku nowym

    os.replace(OUTPUT_PATH, INPUT_PATH)


    data = pd.read_csv(r"C:\Users\Wassup_Home\PycharmProjects\kowalski_analysis\src\car_prices.csv")

    data = data.head()

    doc = Document()
    table = doc.add_table(rows=data.shape[0] + 1, cols=data.shape[1])
    table.style = 'Table Grid'

    for i, columname in enumerate(data.columns):
        table.cell(0, i).text = columname

    for row_idx, row in data.iterrows():
        for col_idx, col in enumerate(row):
            table.cell(row_idx + 1, col_idx).text = str(col)

    docx_path = r"car_prices_head.docx"
    doc.save(docx_path)

def add_age_column():
    # Ścieżka do pliku

    # Wczytaj dane z pliku
    data = pd.read_csv(INPUT_PATH)

    # Aktualny rok
    current_year = datetime.now().year

    # Dodaj nową kolumnę age (wiek)
    data['age'] = current_year - data['year']

    # Zapisz zaktualizowane dane z powrotem do pliku
    data.to_csv(INPUT_PATH, index=False)

def standardize_body():
    # Ścieżka do pliku
    input_path = r"C:\Users\Wassup_Home\PycharmProjects\kowalski_analysis\src\car_prices.csv"

    # Wczytanie danych
    data = pd.read_csv(input_path)

    # Definiujemy mapowanie niestandardowych wartości do ustandaryzowanych
    standardization_map = {
        'suv': 'suv',
        'sedan': 'sedan',
        'convertible': 'convertible',
        'coupe': 'coupe',
        'wagon': 'wagon',
        'hatchback': 'hatchback',
        'pickup': 'pickup',
        'minivan': 'van',
        'van': 'van',
        'e-series van': 'van',
        'transit van': 'van',
        'promaster cargo van': 'van',
        'cts wagon': 'wagon',
        'tsx sport wagon': 'wagon',

        # Zmiana nazw marek na typy nadwozia
        'g sedan': 'sedan',
        'g coupe': 'coupe',
        'g convertible': 'convertible',
        'g37 convertible': 'convertible',
        'q60 convertible': 'convertible',
        'q60 coupe': 'coupe',
        'g37 coupe': 'coupe',
        'genesis coupe': 'coupe',
        'elantra coupe': 'coupe',
        'cts coupe': 'coupe',
        'cts-v coupe': 'coupe',
        'cts-v wagon': 'wagon',
        'beetle convertible': 'convertible',
        'granturismo convertible': 'convertible'
    }


    # Funkcja do standaryzacji wartości
    def standardize_body(value):
        value = str(value).lower()  # Konwersja na małe litery
        return standardization_map.get(value, value)  # Zamiana na ustandaryzowaną wartość


    # Liczba wierszy do przetworzenia
    total_rows = len(data)

    # Pasek postępu + przetwarzanie
    for index in tqdm(data.index, total=total_rows, desc="Standaryzacja danych", unit="wiersz"):
        data.loc[index, 'body'] = standardize_body(data.loc[index, 'body'])

    # Wyświetlenie unikalnych wartości po zmianie
    print("\nUnikalne wartości po standardyzacji:")
    print(data['body'].unique())

    # Zapis do pliku CSV
    data.to_csv(input_path, index=False)
    print("\n✅ Standardyzacja zakończona! Plik zapisano.")
